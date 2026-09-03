# -*- coding: utf-8 -*-
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from apa7_word_components import (
    add_apa_table,
    append_stat_runs,
    apply_apa7_layout,
    format_or,
    format_or_ci,
    format_p,
    format_probability,
)
from edit_apa7_docx import edit_document
from validate_apa7_docx import validate_docx


HERE = Path(__file__).resolve().parent


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


class WordComponentTests(unittest.TestCase):
    def test_statistical_formatters_and_runs_follow_apa(self):
        self.assertEqual(format_probability(0.20904), ".209")
        self.assertEqual(format_p(0.0071), ".007")
        self.assertEqual(format_p(0.0005), "< .001")
        self.assertEqual(format_or(0.6238), "0.62")
        self.assertEqual(format_or_ci(0.6238, 0.5385, 0.7225),
                         "0.62 [0.54, 0.72]")

        doc = Document()
        p = doc.add_paragraph()
        append_stat_runs(
            p,
            "ORs ranged from 0.62 to 2.21, p < .001; n1/(n0 + n1).",
        )
        italic = {r.text for r in p.runs if r.italic}
        self.assertIn("OR", italic)
        self.assertIn("p", italic)
        self.assertIn("n", italic)
        subscript = {r.text for r in p.runs if r.font.subscript}
        self.assertTrue({"0", "1"}.issubset(subscript))

    def test_table_has_apa_borders_widths_and_pagination_controls(self):
        doc = Document()
        apply_apa7_layout(doc)
        table = add_apa_table(
            doc,
            ["Endpoint", "OR [95% CI]", "Holm p"],
            [
                ["Kimi", "0.62 [0.54, 0.72]", "< .001"],
                ["Qwen", "2.21 [1.91, 2.56]", "< .001"],
            ],
            widths=[2.0, 2.7, 1.8],
            font_size=9,
        )
        xml = table._tbl.xml
        self.assertIn('w:type="fixed"', xml)
        self.assertIn("w:tblHeader", table.rows[0]._tr.xml)
        self.assertTrue(all("w:cantSplit" in row._tr.xml for row in table.rows))
        for edge in ("start", "end", "insideV"):
            self.assertRegex(xml, rf"<w:{edge}[^>]+w:val=\"nil\"")
        self.assertEqual(table.cell(1, 0).paragraphs[0].alignment,
                         WD_ALIGN_PARAGRAPH.LEFT)
        self.assertTrue(any(r.text == "OR" and r.italic
                            for r in table.cell(0, 1).paragraphs[0].runs))
        self.assertTrue(any(r.text == "p" and r.italic
                            for r in table.cell(0, 2).paragraphs[0].runs))


class ExistingDocumentEditTests(unittest.TestCase):
    def test_replaces_section_in_copy_and_preserves_header_and_other_sections(self):
        with tempfile.TemporaryDirectory(prefix="apa7_edit_test_") as tmp:
            tmp = Path(tmp)
            source = tmp / "source.docx"
            output = tmp / "edited.docx"
            doc = Document()
            header = doc.sections[0].header.paragraphs[0]
            header.add_run("STUDY TITLE\t")
            add_page_field(header)
            doc.add_heading("Method", level=1)
            doc.add_paragraph("Keep method content.")
            doc.add_heading("Results", level=1)
            doc.add_paragraph("Remove old result.")
            doc.add_heading("Discussion", level=1)
            doc.add_paragraph("Keep discussion content.")
            doc.save(source)
            before = hashlib.sha256(source.read_bytes()).hexdigest()

            spec = {
                "normalize_apa": True,
                "operations": [{
                    "type": "replace_section",
                    "heading": "Results",
                    "content": [
                        {"type": "paragraph",
                         "text": "The contrast was OR = 0.62, 95% CI [0.54, 0.72], p < .001."},
                        {"type": "heading", "level": 3,
                         "text": "Simple Effects"},
                        {"type": "table",
                         "headers": ["Endpoint", "OR [95% CI]", "p"],
                         "rows": [["Kimi", "0.62 [0.54, 0.72]", "< .001"]],
                         "widths": [2.0, 2.7, 1.8]},
                    ],
                }],
            }
            edit_document(source, output, spec)

            self.assertEqual(before, hashlib.sha256(source.read_bytes()).hexdigest())
            edited = Document(output)
            body = "\n".join(p.text for p in edited.paragraphs)
            self.assertIn("Keep method content.", body)
            self.assertIn("Keep discussion content.", body)
            self.assertIn("The contrast was OR = 0.62", body)
            self.assertNotIn("Remove old result.", body)
            self.assertEqual(len(edited.tables), 1)
            inserted_h3 = next(p for p in edited.paragraphs if p.text == "Simple Effects")
            self.assertTrue(all(r.bold and r.italic for r in inserted_h3.runs if r.text))
            self.assertIn("STUDY TITLE", edited.sections[0].header.paragraphs[0].text)
            self.assertIn("PAGE", edited.sections[0].header.paragraphs[0]._p.xml)
            self.assertAlmostEqual(edited.sections[0].left_margin.inches, 1.0)

    def test_validator_accepts_edited_document_and_rejects_unformatted_document(self):
        with tempfile.TemporaryDirectory(prefix="apa7_validate_test_") as tmp:
            tmp = Path(tmp)
            good = tmp / "good.docx"
            bad = tmp / "bad.docx"
            doc = Document()
            apply_apa7_layout(doc)
            add_page_field(doc.sections[0].header.paragraphs[0])
            doc.add_paragraph("A formatted paragraph.")
            doc.save(good)
            Document().save(bad)
            self.assertTrue(validate_docx(good).ok)
            report = validate_docx(bad)
            self.assertFalse(report.ok)
            self.assertTrue(any("font" in e.lower() or "页码" in e for e in report.errors))

    def test_validator_detects_unitalicized_statistical_symbols(self):
        with tempfile.TemporaryDirectory(prefix="apa7_stats_validate_") as tmp:
            path = Path(tmp) / "bad_stats.docx"
            doc = Document()
            apply_apa7_layout(doc)
            add_page_field(doc.sections[0].header.paragraphs[0])
            doc.add_paragraph("The contrast was OR = 0.62, p < .001.")
            doc.save(path)
            report = validate_docx(path)
            self.assertFalse(report.ok)
            self.assertTrue(any("统计符号" in error for error in report.errors))

    def test_validator_does_not_treat_model_version_as_mean_symbol(self):
        with tempfile.TemporaryDirectory(prefix="apa7_model_name_validate_") as tmp:
            path = Path(tmp) / "model_name.docx"
            doc = Document()
            apply_apa7_layout(doc)
            add_page_field(doc.sections[0].header.paragraphs[0])
            doc.add_paragraph("The endpoints included MiniMax M2.7 and Qwen 3.6 Plus.")
            doc.save(path)
            self.assertTrue(validate_docx(path).ok)


class CreateDocumentIntegrationTests(unittest.TestCase):
    def test_create_cli_renders_json_table_spec(self):
        with tempfile.TemporaryDirectory(prefix="apa7_create_test_") as tmp:
            tmp = Path(tmp)
            body = tmp / "body.txt"
            tables = tmp / "tables.json"
            output = tmp / "paper.docx"
            body.write_text(
                "[H1]method\n"
                "The analysis used two stages, OR = 0.62, p < .001; n1/(n0 + n1).\n"
                "[TABLE]\n"
                "[CAPTION]stage 1 language contrasts\n"
                "[TABLEDATA]stage1\n"
                "[NOTE]OR = odds ratio. p values are Holm corrected.\n",
                encoding="utf-8",
            )
            tables.write_text(json.dumps({
                "stage1": {
                    "headers": ["Endpoint", "OR [95% CI]", "Holm p"],
                    "rows": [["Kimi", "0.62 [0.54, 0.72]", "< .001"]],
                    "widths": [2.0, 2.7, 1.8],
                    "font_size": 9,
                }
            }), encoding="utf-8")
            proc = subprocess.run(
                [sys.executable, str(HERE / "create_apa7_docx.py"),
                 "--title", "language effects in large language models",
                 "--body", str(body), "--tables", str(tables),
                 "--output", str(output)],
                text=True, capture_output=True, encoding="utf-8",
            )
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            doc = Document(output)
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(doc.tables[0].cell(1, 0).text, "Kimi")
            self.assertIn("w:tblHeader", doc.tables[0].rows[0]._tr.xml)
            stats = next(p for p in doc.paragraphs if "analysis used two stages" in p.text)
            self.assertTrue(any(r.text == "OR" and r.italic for r in stats.runs))
            self.assertTrue(any(r.text == "p" and r.italic for r in stats.runs))
            self.assertTrue(any(r.text == "1" and r.font.subscript for r in stats.runs))


if __name__ == "__main__":
    unittest.main(verbosity=2)
