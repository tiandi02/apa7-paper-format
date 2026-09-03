# -*- coding: utf-8 -*-
"""独立检查 APA 7 DOCX 的结构、版式、表格和统计格式。"""

import argparse
from dataclasses import dataclass, field
import json
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile

from docx import Document
from docx.oxml.ns import qn

from apa7_word_components import FONT, STAT_RE


@dataclass
class ValidationReport:
    path: str
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    metrics: dict = field(default_factory=dict)

    @property
    def ok(self):
        return not self.errors

    def as_dict(self):
        return {
            "path": self.path,
            "ok": self.ok,
            "errors": self.errors,
            "warnings": self.warnings,
            "metrics": self.metrics,
        }


def _has_on_off(element, name):
    return element.find(qn(f"w:{name}")) is not None


def _paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _table_errors(table, index, content_width_dxa):
    errors = []
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None or layout.get(qn("w:type")) != "fixed":
        errors.append(f"表格 {index} 未使用固定布局")
    borders = tbl_pr.find(qn("w:tblBorders"))
    for edge in ("start", "end", "insideV"):
        elm = borders.find(qn(f"w:{edge}")) if borders is not None else None
        if elm is None or elm.get(qn("w:val")) != "nil":
            errors.append(f"表格 {index} 含不符合 APA 的竖线：{edge}")
    if not table.rows or not _has_on_off(table.rows[0]._tr.get_or_add_trPr(),
                                        "tblHeader"):
        errors.append(f"表格 {index} 未设置重复表头")
    for row_index, row in enumerate(table.rows, 1):
        if not _has_on_off(row._tr.get_or_add_trPr(), "cantSplit"):
            errors.append(f"表格 {index} 第 {row_index} 行允许跨页拆分")
    grid = table._tbl.tblGrid
    widths = [int(col.get(qn("w:w"), "0")) for col in grid.gridCol_lst]
    if widths and sum(widths) > content_width_dxa + 5:
        errors.append(f"表格 {index} 列宽超过正文宽度")
    return errors


def _stat_symbol_errors(doc):
    errors = []
    for paragraph in _paragraphs(doc):
        spans = []
        offset = 0
        for run in paragraph.runs:
            spans.append((offset, offset + len(run.text), run))
            offset += len(run.text)
        for match in STAT_RE.finditer(paragraph.text):
            runs = [run for start, end, run in spans
                    if start <= match.start() < end]
            if not runs:
                continue
            run = runs[0]
            inherited = paragraph.style.font.italic if paragraph.style else None
            if run.italic is not True and inherited is not True:
                context = paragraph.text.strip()[:80]
                errors.append(
                    f"统计符号 {match.group(0)} 未使用斜体：{context}")
    return errors


def validate_docx(path):
    path = Path(path)
    report = ValidationReport(str(path))
    if not path.exists():
        report.errors.append("文件不存在")
        return report
    try:
        with ZipFile(path) as archive:
            bad_member = archive.testzip()
            if bad_member:
                report.errors.append(f"ZIP CRC 错误：{bad_member}")
            document_xml = archive.read("word/document.xml").decode("utf-8")
            if re.search(r"<w:(?:ins|del)(?:\s|>)", document_xml):
                report.warnings.append("文档包含修订标记；请确认是否需要接受修订")
    except (BadZipFile, KeyError, UnicodeDecodeError) as exc:
        report.errors.append(f"DOCX 包结构无效：{exc}")
        return report

    try:
        doc = Document(path)
    except Exception as exc:
        report.errors.append(f"python-docx 无法打开文档：{exc}")
        return report

    report.metrics.update({
        "sections": len(doc.sections),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
    })
    normal = doc.styles["Normal"]
    if normal.font.name != FONT or normal.font.size is None or abs(normal.font.size.pt - 12) > .01:
        report.errors.append("Normal 样式字体应为 12 pt Times New Roman")
    if normal.paragraph_format.line_spacing != 2.0:
        report.errors.append("Normal 样式应为双倍行距")
    if normal.paragraph_format.first_line_indent is None or abs(normal.paragraph_format.first_line_indent.inches - .5) > .001:
        report.errors.append("Normal 样式首行缩进应为 0.5 英寸")
    if normal.paragraph_format.widow_control is not True:
        report.errors.append("Normal 样式未开启孤行寡行控制")

    for index, section in enumerate(doc.sections, 1):
        if abs(section.page_width.inches - 8.5) > .001 or abs(section.page_height.inches - 11) > .001:
            report.errors.append(f"第 {index} 节纸张不是 US Letter")
        margins = [section.top_margin, section.bottom_margin,
                   section.left_margin, section.right_margin]
        if any(m is None or abs(m.inches - 1) > .001 for m in margins):
            report.errors.append(f"第 {index} 节页边距不是 1 英寸")
        if "PAGE" not in section.header.paragraphs[0]._p.xml:
            report.errors.append(f"第 {index} 节页眉缺少自动页码")

    if doc.sections:
        section = doc.sections[-1]
        content_width_dxa = round(
            (section.page_width - section.left_margin - section.right_margin)
            / 914400 * 1440)
        for index, table in enumerate(doc.tables, 1):
            report.errors.extend(_table_errors(table, index, content_width_dxa))

    text = "\n".join(p.text for p in _paragraphs(doc))
    if re.search(r"\bOR\s*(?:=|was)\s*\.\d", text):
        report.errors.append("OR 小于 1 时缺少前导零")
    if re.search(r"\bOR\s*=\s*\d+(?:\.\d+)?\s*,?\s*95% CI \[\.\d", text):
        report.errors.append("OR 的置信区间下限缺少前导零")
    report.errors.extend(_stat_symbol_errors(doc))
    return report


def main():
    parser = argparse.ArgumentParser(description="验证 APA 7 DOCX")
    parser.add_argument("path", help="待验证 DOCX")
    parser.add_argument("--json", action="store_true", help="以 JSON 输出报告")
    args = parser.parse_args()
    report = validate_docx(args.path)
    if args.json:
        print(json.dumps(report.as_dict(), ensure_ascii=False, indent=2))
    else:
        print("PASS" if report.ok else "FAIL", report.path)
        for error in report.errors:
            print("错误：", error)
        for warning in report.warnings:
            print("警告：", warning)
    raise SystemExit(0 if report.ok else 1)


if __name__ == "__main__":
    main()
