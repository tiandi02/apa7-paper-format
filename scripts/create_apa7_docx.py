# -*- coding: utf-8 -*-
"""生成符合 APA 7 版式的论文 .docx（学生版/专业版）。

版式：1 英寸页边距、Times New Roman 12pt（中文宋体）、全文双倍行距、
每页页眉含页码（右上角；专业版为左对齐 running head + 右对齐页码同一行）、
标题页居中、参考文献悬挂缩进、孤行寡行控制（widowControl）。
生成后自动读回校验版式，失败即报错。

用法示例：
  # 学生论文（含摘要页；多作者/多单位用 ; 分隔，^1^ 上标与单位编号对应）
  python create_apa7_docx.py --title "the effect of X on Y" \\
      --author "Ming Li^1^; Zhang San^2^" \\
      --affiliation "Department of Psychology, XX University^1^; School of Education, YY University^2^" \\
      --course "PSY204" --instructor "Dr. Zhang" --date "March 16, 2026" \\
      --abstract abstract.txt --keywords "kw1, kw2" --keywords-protect "COVID-19, Big Five" \\
      --body body.txt --references refs.txt --output paper.docx
  # 专业论文（投稿：作者注 + running head）
  python create_apa7_docx.py --mode professional --title "..." --author "..." \\
      --author-note note.txt --running-head "..." --output paper.docx
  # 回归自检（固定样例 + 版式断言）
  python create_apa7_docx.py --selftest

正文文件（--body）标记语法：
  [H1]-[H5]   行首：五级标题。英文标题自动转 title case（6.17 规则），中文跳过。
               H4/H5 标题以句点结尾、正文同行：按第一个 ". " 拆分，仅标题部分加粗（H5 加粗斜体）。
  [QUOTE]     进入块引用模式（整段左缩进 0.5 英寸、首行不缩进）
  [P]         块引用内的新段落（首行再缩进 0.5 英寸）
  [ENDQUOTE]  退出块引用模式
  [TABLE]     表格编号行（不写编号则自动 Table 1、2……；附录内自动 Table A1、A2……），加粗、左对齐
  [CAPTION]   表题/图题（斜体、title case、左对齐）；紧跟 [APPENDIX] 之后则为附录标题（加粗、居中）
  [NOTE]      表格/图注（"Note." 斜体开头；特定注的字母用 ^a^ 上标标记）
  [TABLEDATA] 由 --tables 指定的 JSON 表格 ID；在当前位置插入真正的 APA Word 表格
  [FIGURE]    图编号行（自动 Figure 1、2……；附录内 Figure C2），加粗、左对齐
  [APPENDIX]  附录标签行（如 [APPENDIX]Appendix A），加粗、居中，并分页；
              其后 [TABLE]/[FIGURE] 按该附录字母自动编号，直到出现正文段落或标题
  [EQ]        显示公式行：真正的 Word 公式对象（OMML）、居中、编号右对齐 (1)(2)……
              支持的公式方言：变量斜体、_ 下标、^ 上标、\\frac{a}{b} 分式、\\sqrt{x}/\\sqrt[3]{x} 根式、
              \\sum_{i=1}^{n} 上下限、希腊字母（\\alpha \\beta \\Sigma 等）、函数名正体（log exp sin 等）
  行内标记：*斜体*、^上标^、~下标~、{FN:脚注内容}（真实 Word 脚注，位于页脚单倍行距）、
            {REF:Table 1}（题注交叉引用域）、{EQREF:1}（公式交叉引用域）、
            URL/DOI 自动转真实 Word 超链接
  --author/--affiliation 等参数同样支持行内标记（如作者右上标编号 "Ming Li^1^"）

其他行为：
  - 参考文献默认按 APA 9.44–9.49 完整规则自动排序（单作者优先、同作者按日期 n.d. 最前 /
    in press 最后、同第一作者按后续作者、无作者按标题、nothing precedes something），
    并为同第一作者同年的条目自动加 a/b/c 后缀（--no-sort 可关闭）
  - 生成前粗查文内年份与参考文献年份的对应关系，不对应时输出警告
  - 摘要超过 250 词、running head 超过 50 字符时输出警告
  - 关键词除专有名词外自动转小写（2.10）：专名信号（全大写、含数字、词中大写、
    非首词首字母大写）自动保留，--keywords-protect 可显式指定白名单
"""
import argparse
import contextlib
import html
import io
import json
import os
import re
import sys
import tempfile
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.package import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from apa7_word_components import add_apa_table_spec, iter_stat_segments

# Windows 控制台默认 GBK，重配为 UTF-8 避免中文输出乱码
sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

FONT = "Times New Roman"
FONT_EAST = "宋体"
MATH_FONT = "Cambria Math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
INLINE_RE = re.compile(
    r"(\*[^*\n]+\*|\^[^\^\n]+\^|~[^~\n]+~"
    r"|\{FN:[^}\n]+\}|\{EQREF:[^}\n]+\}|\{REF:[^}\n]+\}"
    r"|https?://[^\s]+)")
MINOR_WORDS = {"a", "an", "the", "and", "but", "or", "nor", "for", "so", "yet",
               "if", "at", "by", "in", "of", "on", "to", "up", "as", "per", "via", "off"}
# 6.17：冒号/破折号/句末标点后的首词大写（句点仅限后跟空白，避免拆开 U.S. 之类缩写）
SEP_RE = re.compile(r"(:|—|–|!|\?|\.(?=\s))")

_esc = html.escape


def title_case(s):
    """按 APA 6.17 转 title case：首词及冒号/破折号/句末标点后的首词大写、
    ≥4 字母的词大写、连字符复合词两部分均大写、短冠词/连词/介词小写；
    单字母词保留原样（Appendix A、Factor B）；含中文的原样返回。"""
    if not s or re.search(r"[一-鿿]", s):
        return s

    def convert(word):
        if "-" in word:
            return "-".join(x.capitalize() if x.islower() else x
                            for x in word.split("-"))
        return word.capitalize() if word.islower() else word

    def convert_part(part):
        lead = part[:len(part) - len(part.lstrip())]  # 保留段首空白（标点后的空格）
        words = part.split()
        out = []
        for i, w in enumerate(words):
            if i == 0:
                out.append(convert(w))
            elif w.lower() in MINOR_WORDS:
                # 单字母词是标签（Appendix A、Factor B），大写；其余短词小写
                out.append(w.upper() if len(w) == 1 else w.lower())
            else:
                out.append(convert(w))
        return lead + " ".join(out)

    return "".join(
        p if p in (":", "—", "–", "!", "?", ".") else convert_part(p)
        for p in SEP_RE.split(s) if p)


def set_normal_style(doc):
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST)  # 中文用宋体
    pf = style.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    # 孤行寡行控制（widowControl），写入 Normal 样式
    pPr = style.element.get_or_add_pPr()
    if pPr.find(qn("w:widowControl")) is None:
        pPr.insert(0, OxmlElement("w:widowControl"))


def set_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def _set_run_font(run):
    run.font.name = FONT
    run.font.size = Pt(12)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST)


def _add_page_field(p):
    """向段落追加 PAGE 域（自动页码）。"""
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    _set_run_font(run)


def setup_header(doc, running_head=""):
    """遍历全部 section 写页眉：学生版仅右对齐页码；专业版左 running head 与
    右对齐页码在同一行（用右制表位实现，不排两行）。"""
    for section in doc.sections:
        p = section.header.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)  # 清空默认段落
        if running_head:
            r = p.add_run(running_head)
            _set_run_font(r)
            p.paragraph_format.tab_stops.add_tab_stop(
                Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run("\t")
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _add_page_field(p)


def _make_running_head(title, given):
    """专业论文 running head（2.8）：缺省由标题生成——大写、& 代替 and、截 50 字符。"""
    if given:
        return given
    return title.replace(" and ", " & ").upper()[:50]


class DocState:
    """跨段落状态：脚注部件与 id、公式编号、书签 id（全文档唯一）、
    题注查重、附录字母与表图计数器、前向引用的已知书签名。"""

    def __init__(self):
        self.footnotes_part = None
        self.footnote_id = 1
        self.eq_no = 0
        self.bm_id = 0
        self.bookmarks = {}        # 已创建书签名 -> 重名次数
        self.known_bookmarks = set()  # 全部编号（含前向引用，预扫描登记）
        self.appendix = None
        self.app_counters = {}
        self.table_no = 0
        self.figure_no = 0
        self.table_specs = {}


def _make_run_elm(text, bold=False, italic=False, sup=False, sub=False):
    """构造带 TNR 12pt 字体的 w:r 元素（正文/脚注/超链接内的 run 通用）。"""
    parts = [f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT_EAST}"/>',
             '<w:sz w:val="24"/><w:szCs w:val="24"/>']
    if bold:
        parts.append('<w:b/>')
    if italic:
        parts.append('<w:i/>')
    if sup:
        parts.append('<w:vertAlign w:val="superscript"/>')
    if sub:
        parts.append('<w:vertAlign w:val="subscript"/>')
    return parse_xml(f'<w:r xmlns:w="{W_NS}"><w:rPr>{"".join(parts)}</w:rPr>'
                     f'<w:t xml:space="preserve">{_esc(text)}</w:t></w:r>')


def _add_settings_fields(doc):
    """settings.xml 补 w:updateFields（打开时自动更新域）与 w:footnotePr（脚注设置）。
    两者按 schema 序插到 w:compat 之前，幂等。"""
    settings = doc.settings.element
    anchor = next((c for c in settings if c.tag in {
        qn("w:compat"), qn("w:rsids"), qn("w:mathPr"), qn("w:themeFontLang"),
        qn("w:clrSchemeMapping"), qn("w:shapeDefaults")}), None)

    def insert(elm):
        if anchor is not None:
            anchor.addprevious(elm)
        else:
            settings.append(elm)

    if settings.find(qn("w:updateFields")) is None:
        insert(parse_xml(f'<w:updateFields xmlns:w="{W_NS}" w:val="true"/>'))
    if settings.find(qn("w:footnotePr")) is None:
        insert(parse_xml(
            f'<w:footnotePr xmlns:w="{W_NS}">'
            '<w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotePr>'))


FOOTNOTES_TEMPLATE = (
    f'<w:footnotes xmlns:w="{W_NS}" xmlns:r="{R_NS}">'
    '<w:footnote w:type="separator" w:id="-1">'
    '<w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
    '<w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0">'
    '<w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
    '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    '</w:footnotes>')


def _ensure_footnotes_part(doc, state):
    """惰性创建 /word/footnotes.xml 部件并挂到文档关系（不 relate 的部件不会被写入包）。"""
    if state.footnotes_part is None:
        root = parse_xml(FOOTNOTES_TEMPLATE)
        state.footnotes_part = XmlPart(PackURI("/word/footnotes.xml"),
                                       CT.WML_FOOTNOTES, root, doc.part.package)
        doc.part.relate_to(state.footnotes_part, RT.FOOTNOTES)
        _add_settings_fields(doc)
    return state.footnotes_part


def _add_footnote_paragraph(doc, state, n, text):
    """向脚注部件追加一条内容脚注（单倍行距、首行不缩进、footnoteRef 自动编号）。"""
    fn = parse_xml(
        f'<w:footnote xmlns:w="{W_NS}" w:id="{n}">'
        '<w:p><w:pPr>'
        '<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        '<w:ind w:firstLine="0"/>'
        '</w:pPr>'
        '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
        f'<w:r><w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT_EAST}"/>'
        '<w:sz w:val="24"/><w:szCs w:val="24"/></w:rPr>'
        '<w:t xml:space="preserve"> </w:t></w:r>'
        '</w:p></w:footnote>')
    p_elm = fn[0]
    # 脚注内的超链接关系必须挂到脚注部件自己的 rels
    append_runs_to_p_elm(p_elm, text, doc=doc, state=state,
                         owner_part=state.footnotes_part)
    state.footnotes_part.element.append(fn)


def _footnote_ref_xml(n):
    """正文内的脚注引用 run（上标数字，Word 打开时自动编号）。"""
    return parse_xml(
        f'<w:r xmlns:w="{W_NS}"><w:rPr>'
        f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT_EAST}"/>'
        '<w:sz w:val="24"/><w:szCs w:val="24"/>'
        '<w:vertAlign w:val="superscript"/>'
        f'</w:rPr><w:footnoteReference w:id="{n}"/></w:r>')


def _add_hyperlink(p_elm, url, owner_part):
    """URL/DOI 转真实 Word 超链接（APA 9.35：电子阅读需为可点击链接）。
    r:id 关系必须建在元素所在部件（正文 doc.part、脚注 footnotes_part）。"""
    rId = owner_part.relate_to(url, RT.HYPERLINK, is_external=True)
    p_elm.append(parse_xml(
        f'<w:hyperlink xmlns:w="{W_NS}" xmlns:r="{R_NS}" r:id="{rId}">'
        f'<w:r><w:rPr>'
        f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT_EAST}"/>'
        '<w:sz w:val="24"/><w:szCs w:val="24"/>'
        f'</w:rPr><w:t xml:space="preserve">{_esc(url)}</w:t></w:r></w:hyperlink>'))


def _add_ref_field(doc, p_elm, bookmark_name, placeholder_text=None):
    """向段落追加 REF 域（\\h 开关 = 指向书签的超链接）。separate 与 end 之间
    放占位结果文本：域未更新时（如刚打开）该处仍可见，不显示空白。"""
    if doc is not None:
        _add_settings_fields(doc)
    xml = (f'<w:r xmlns:w="{W_NS}"><w:rPr>'
           f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT_EAST}"/>'
           '<w:sz w:val="24"/><w:szCs w:val="24"/></w:rPr>'
           '<w:fldChar w:fldCharType="begin"/>'
           f'<w:instrText xml:space="preserve"> REF {bookmark_name} \\h </w:instrText>'
           '<w:fldChar w:fldCharType="separate"/>')
    if placeholder_text:
        xml += f'<w:t xml:space="preserve">{_esc(placeholder_text)}</w:t>'
    xml += '<w:fldChar w:fldCharType="end"/></w:r>'
    p_elm.append(parse_xml(xml))


def append_runs_to_p_elm(p_elm, text, base_bold=False, base_italic=False,
                         tc=False, doc=None, state=None, owner_part=None):
    """向 w:p 元素追加 run 序列。识别行内标记 * ^ ~、脚注 {FN:}、交叉引用
    {REF:}/{EQREF:} 与 URL 超链接；owner_part 决定超链接关系挂到哪个部件。"""
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        if seg.startswith("http"):
            url = re.sub(r"[.,;:)\]]+$", "", seg)  # 剥离句尾标点（DOI 后不加句点）
            if owner_part is None:
                p_elm.append(_make_run_elm(seg))
            else:
                _add_hyperlink(p_elm, url, owner_part)
            continue
        if seg.startswith("{FN:") and seg.endswith("}"):
            if doc is not None and state is not None:
                _ensure_footnotes_part(doc, state)
                n = state.footnote_id
                state.footnote_id += 1
                _add_footnote_paragraph(doc, state, n, seg[4:-1])
                p_elm.append(_footnote_ref_xml(n))
            else:
                print("警告：{FN:} 标记需要文档上下文，已按普通文本输出：", seg[4:-1])
                p_elm.append(_make_run_elm(seg[4:-1]))
            continue
        if seg.startswith("{EQREF:") and seg.endswith("}"):
            n = seg[7:-1].strip()
            if state is not None:
                _add_ref_field(doc, p_elm, f"Eq{n}", f"({n})")
            else:
                p_elm.append(_make_run_elm(f"({n})"))
            continue
        if seg.startswith("{REF:") and seg.endswith("}"):
            target = seg[5:-1].strip()
            name = target.replace(" ", "_")
            if state is not None and name in state.known_bookmarks:
                _add_ref_field(doc, p_elm, name, target)
            else:
                if state is not None:
                    print(f"警告：交叉引用目标不存在：{target}，已按普通文本输出")
                p_elm.append(_make_run_elm(target))
            continue
        bold, italic, sup, sub = base_bold, base_italic, False, False
        s = seg
        if seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
            s, italic = seg[1:-1], True
        elif seg.startswith("^") and seg.endswith("^") and len(seg) > 2:
            s, sup = seg[1:-1], True
        elif seg.startswith("~") and seg.endswith("~") and len(seg) > 2:
            s, sub = seg[1:-1], True
        if tc:
            s = title_case(s)
        if sup or sub:
            p_elm.append(_make_run_elm(s, bold, italic, sup, sub))
        else:
            for value, stat_italic, stat_sub in iter_stat_segments(s):
                p_elm.append(_make_run_elm(
                    value, bold, italic or stat_italic, sub=stat_sub))


def add_runs_with_markup(p, text, base_bold=False, base_italic=False, tc=False,
                         doc=None, state=None):
    """兼容适配器：向 python-docx 段落追加带标记的 run 序列。"""
    append_runs_to_p_elm(p._p, text, base_bold, base_italic, tc, doc, state,
                         owner_part=doc.part if doc is not None else None)


def para(doc, text="", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False,
         indent=None, left=None, hanging=False, tc=False, state=None):
    p = doc.add_paragraph()
    p.alignment = align
    if indent is not None:
        p.paragraph_format.first_line_indent = Inches(indent)
    if left is not None:
        p.paragraph_format.left_indent = Inches(left)
    if hanging:
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
    if text:
        add_runs_with_markup(p, text, bold, italic, tc, doc, state)
    return p


def _join_authors(authors):
    """APA 2.6：2 位作者 and 连接；3 位及以上逗号分隔、末位前加 and。"""
    if len(authors) == 1:
        return authors[0]
    if len(authors) == 2:
        return f"{authors[0]} and {authors[1]}"
    return ", ".join(authors[:-1]) + f", and {authors[-1]}"


def title_page(doc, args, title):
    for _ in range(3):  # 标题位于页面上半部（距顶约 3-4 行）
        para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)  # 标题与署名之间隔一个双倍空行
    authors = [a.strip() for a in args.author.split(";") if a.strip()]
    if authors:
        para(doc, _join_authors(authors), align=WD_ALIGN_PARAGRAPH.CENTER)
    affiliations = [a.strip() for a in args.affiliation.split(";") if a.strip()]
    for aff in affiliations:  # 多单位各占一行，^1^ 上标编号与作者对应（2.6）
        para(doc, aff, align=WD_ALIGN_PARAGRAPH.CENTER)
    if args.mode == "professional":
        if args.author_note:
            note_lines = open(args.author_note, encoding="utf-8").readlines()
            para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
            para(doc, "Author Note", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            for line in note_lines:
                line = line.strip()
                if line:
                    para(doc, line, indent=0.5)  # 2.7：正文左对齐、首行缩进 0.5
    else:
        for line in (args.course, args.instructor, args.date):
            if line:
                para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER)


def lowercase_keywords(text, protect=""):
    """关键词除专有名词外转小写（2.10）。返回 (处理后的关键词串, 保留词列表)。

    保留规则：--keywords-protect 白名单整条保留；词级专名信号（全大写、含数字且非
    全小写如 COVID-19、词中大写如 iPhone、非首词首字母大写）保留；首词若同关键词
    已有其他保留词且原本首字母大写，则一并保留（保守避免破坏 United States、
    Big Five 这类专名短语——误保留可由用户人工降为小写，误破坏则改不回原专名）。"""
    protected = {p.strip().lower() for p in protect.split(",") if p.strip()}
    out = []
    kept = []

    def looks_proper(w):
        return (len(w) >= 2 and w.isupper()
                or re.search(r"[A-Za-z]", w) and re.search(r"\d", w) and not w.islower()
                or re.search(r"[a-z][A-Z]", w))

    for kw in text.split(","):
        kw = kw.strip()
        if not kw:
            continue
        if kw.lower() in protected:
            out.append(kw)
            kept.append(kw)
            continue
        words = kw.split()
        res = []
        rest_kept = False  # 该关键词内除首词外是否有保留词
        for i, w in enumerate(words):
            if looks_proper(w) or (i > 0 and w[0].isupper()):
                res.append(w)
                kept.append(w)
                if i > 0:
                    rest_kept = True
            else:
                res.append(w.lower())
        if (rest_kept and res[0] != words[0]
                and words[0] and words[0][0].isupper()):
            res[0] = words[0]
            kept.append(words[0])
        out.append(" ".join(res))
    return ", ".join(out), list(dict.fromkeys(kept))


def abstract_page(doc, args, state=None):
    doc.add_page_break()
    para(doc, "Abstract", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    lines = open(args.abstract, encoding="utf-8").readlines() if args.abstract else []
    text = " ".join(l.strip() for l in lines if l.strip())
    if text:
        n = len(text.split())
        if n > 250:
            print(f"警告：摘要 {n} 词，超过 APA 上限 250 词（2.9），请精简")
        para(doc, text, indent=0, state=state)  # 单段、首行不缩进
    if args.keywords:
        kws, kept = lowercase_keywords(args.keywords,
                                       getattr(args, "keywords_protect", ""))
        if kept:
            print("提示：以下关键词/词保持原大小写（视为专有名词），请核对：",
                  ", ".join(kept),
                  "（可用 --keywords-protect 显式指定白名单）")
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        r = p.add_run("Keywords: ")
        r.italic = True
        _set_run_font(r)
        add_runs_with_markup(p, kws, doc=doc, state=state)


HEADING_STYLES = {
    "H1": dict(align=WD_ALIGN_PARAGRAPH.CENTER, bold=True),
    "H2": dict(align=WD_ALIGN_PARAGRAPH.LEFT, bold=True),
    "H3": dict(align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, italic=True),
}


def render_heading(doc, line):
    """渲染标题行。[H1]-[H3] 整段加粗/斜体；[H4]/[H5] 标题句点结尾、正文同行，
    按第一个 ". " 拆成两个 run——仅标题部分加粗（H5 加粗斜体），正文不加粗。"""
    level = line[1:3]
    text = line.split("]", 1)[1].strip()
    if level in ("H4", "H5"):
        idx = text.find(". ")
        if idx == -1:
            head = text if text.endswith(".") else text + "."
            body_text = ""
        else:
            head, body_text = text[:idx + 1], " " + text[idx + 2:]  # 保留句点后的空格
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        add_runs_with_markup(p, head, base_bold=True,
                             base_italic=(level == "H5"), tc=True)
        if body_text:
            add_runs_with_markup(p, body_text)
    else:
        para(doc, text, **HEADING_STYLES[level], tc=True)


# ---------------- OMML 公式（[EQ] 行） ----------------

TOKEN_RE = re.compile(r"""
    \\[A-Za-z]+
  | [A-Za-z]+
  | \d+(?:\.\d+)?
  | _|\^|\{|\}|\[|\]
  | [+\-−=<>≤≥×÷/(),.]
  | \s+
""", re.X)

GREEK = {
    "alpha": ("α", True), "beta": ("β", True), "gamma": ("γ", True),
    "delta": ("δ", True), "epsilon": ("ε", True), "theta": ("θ", True),
    "lambda": ("λ", True), "mu": ("μ", True), "pi": ("π", True),
    "sigma": ("σ", True), "phi": ("φ", True), "chi": ("χ", True),
    "omega": ("ω", True), "rho": ("ρ", True), "eta": ("η", True),
    "Delta": ("Δ", False), "Sigma": ("Σ", False), "Pi": ("Π", False),
    "Omega": ("Ω", False), "Theta": ("Θ", False),
}
FUNC_NAMES = {"log", "ln", "exp", "sin", "cos", "tan", "max", "min", "det", "lim"}
NARY_CHARS = {"sum": "∑", "prod": "∏", "int": "∫"}


def _m_run(text, italic=False, func=False):
    """OMML 数学 run：文本在 m:t；斜体用 m:sty=i + w:i 双保险（LibreOffice 尊重 w:i）；
    函数名用 m:nor 禁止数学自动斜体；字体 Cambria Math。"""
    nor = "<m:nor/>" if func else ""
    wi = "<w:i/>" if italic else ""
    return (f'<m:r><m:rPr>{nor}<m:sty m:val="{"i" if italic else "p"}"/></m:rPr>'
            f'<w:rPr xmlns:w="{W_NS}">'
            f'<w:rFonts w:ascii="{MATH_FONT}" w:hAnsi="{MATH_FONT}"/>{wi}'
            f'</w:rPr>'
            f'<m:t xml:space="preserve">{_esc(text)}</m:t></m:r>')


def _frac(num, den):
    return (f'<m:f><m:fPr><m:type m:val="bar"/></m:fPr>'
            f'<m:num>{num}</m:num><m:den>{den}</m:den></m:f>')


def _rad(deg, e):
    hide = '<m:degHide m:val="1"/>' if not deg else ""
    d = f'<m:deg>{deg}</m:deg>' if deg else '<m:deg/>'
    return f'<m:rad><m:radPr>{hide}</m:radPr>{d}<m:e>{e}</m:e></m:rad>'


def _sub(base, sub):
    return f'<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>'


def _sup(base, sup):
    return f'<m:sSup><m:e>{base}</m:e><m:sup>{sup}</m:sup></m:sSup>'


def _subsup(base, sub, sup):
    return (f'<m:sSubSup><m:e>{base}</m:e><m:sub>{sub}</m:sub>'
            f'<m:sup>{sup}</m:sup></m:sSubSup>')


def _nary(ch, sub, sup, e):
    pr = (f'<m:naryPr><m:chr m:val="{ch}"/>'
          f'<m:limLoc m:val="{"undOvr" if (sub or sup) else "subSup"}"/>'
          f'<m:subHide m:val="{"0" if sub else "1"}"/>'
          f'<m:supHide m:val="{"0" if sup else "1"}"/></m:naryPr>')
    s = f'<m:sub>{sub}</m:sub>' if sub else ""
    p = f'<m:sup>{sup}</m:sup>' if sup else ""
    return f'<m:nary>{pr}{s}{p}<m:e>{e}</m:e></m:nary>'


def _tokenize(src):
    toks = []
    for m in TOKEN_RE.finditer(src):
        s = m.group(0)
        if s.isspace():
            continue
        if s.startswith("\\") and len(s) > 1:
            toks.append(("cmd", s[1:]))
        elif s in ("_", "^", "{", "}", "[", "]"):
            toks.append((s, s))
        elif s[0].isalpha():
            toks.append(("word", s))
        elif s[0].isdigit() or s[0] == ".":
            toks.append(("num", s))
        else:
            toks.append(("op", s))
    return toks


def _parse_formula(src):
    """递归下降解析公式方言 → m:oMath 内部 XML 字符串。失败抛 ValueError，
    由调用方降级为斜体文本输出（永不中断生成）。"""
    toks = _tokenize(src)
    if not toks:
        raise ValueError("空公式")
    pos = [0]

    def peek():
        return toks[pos[0]] if pos[0] < len(toks) else (None, "")

    def take():
        t = toks[pos[0]]
        pos[0] += 1
        return t

    def parse_expr():
        parts = [parse_term()]
        while peek()[0] == "op" and peek()[1] not in "()":  # 括号由 parse_atom 成组处理
            v = take()[1]
            parts.append(_m_run("−" if v == "-" else v))
            parts.append(parse_term())
        return "".join(parts)

    def parse_term():
        base = parse_atom()
        if peek()[0] == "_":
            take()
            sub = parse_group_or_atom()
            if peek()[0] == "^":
                take()
                sup = parse_group_or_atom()
                return _subsup(base, sub, sup)
            return _sub(base, sub)
        if peek()[0] == "^":
            take()
            sup = parse_group_or_atom()
            return _sup(base, sup)
        return base

    def parse_group_or_atom():
        if peek()[0] == "{":
            return parse_group()
        return parse_atom()

    def parse_group():
        take()  # {
        inner = parse_expr()
        if peek()[0] == "}":
            take()
        return inner  # 花括号只作分组，不输出可见括号

    def parse_atom():
        k, v = take()
        if k is None:
            raise ValueError("公式意外结束")
        if k == "cmd":
            if v == "frac":
                return _frac(parse_group_or_atom(), parse_group_or_atom())
            if v == "sqrt":
                deg = ""
                if peek()[0] == "[":
                    take()
                    deg = parse_expr()
                    if peek()[0] == "]":
                        take()
                return _rad(deg, parse_group_or_atom())
            if v in NARY_CHARS:
                sub = sup = ""
                if peek()[0] == "_":
                    take()
                    sub = parse_group_or_atom()
                if peek()[0] == "^":
                    take()
                    sup = parse_group_or_atom()
                return _nary(NARY_CHARS[v], sub, sup, parse_group_or_atom())
            if v in GREEK:
                ch, it = GREEK[v]
                return _m_run(ch, italic=it)
            if v.lower() in FUNC_NAMES:
                return _m_run(v, func=True)
            return _m_run("\\" + v)  # 未知命令按正体字面输出
        if k == "word":
            if v.lower() in FUNC_NAMES:
                return _m_run(v, func=True)
            if len(v) == 1:
                return _m_run(v, italic=True)
            return _m_run(v)  # 多字母标识符正体（APA 不斜）
        if k == "op" and v == "(":
            # 可见括号成组为原子，使其后 ^ 上标作用于整组：((x_i - M))^2
            inner = parse_expr()
            if peek()[1] == ")":
                take()
            return _m_run("(") + inner + _m_run(")")
        if k in ("num", "op"):
            return _m_run("−" if v == "-" else v)
        if k in ("{", "}", "[", "]"):
            return _m_run(v)  # 不配对的括号按字面输出
        raise ValueError(f"无法解析的记号：{v!r}")

    result = parse_expr()
    if pos[0] < len(toks):
        raise ValueError(f"公式存在无法解析的剩余内容：{toks[pos[0]:]}")
    return result


def render_equation(doc, state, line, number):
    """[EQ] 行：段落制表位居中公式 + 右对齐编号 (n)，编号加 EqN 书签供 {EQREF:n}。"""
    src = line.split("]", 1)[1].strip()
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    try:
        omath = parse_xml(f'<m:oMath xmlns:m="{M_NS}">{_parse_formula(src)}</m:oMath>')
        p._p.append(omath)
    except Exception as e:
        print(f"警告：公式解析失败（{e}），已按斜体文本输出：{src}")
        add_runs_with_markup(p, f"*{src}*", doc=doc, state=state)
    p.add_run("\t")
    r = p.add_run(f"({number})")
    _set_run_font(r)
    _add_bookmark(state, r._r, f"Eq{number}")


def _add_bookmark(state, r_elm, name):
    """在 run 前后插成对书签（w:id 全文档唯一）。书签为 w:p 直接子级。"""
    bid = str(state.bm_id)
    state.bm_id += 1
    s = OxmlElement("w:bookmarkStart")
    s.set(qn("w:id"), bid)
    s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd")
    e.set(qn("w:id"), bid)
    r_elm.addprevious(s)
    r_elm.addnext(e)
    state.bookmarks.setdefault(name, 0)


def _label_para(doc, state, label, bold=False):
    """题注/图注编号行：编号 run 加书签（供 {REF:} 交叉引用），重名加后缀。"""
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = bold
    _set_run_font(r)
    name = label.replace(" ", "_")  # 书签名不能含空格："Table 1" -> "Table_1"
    if name in state.bookmarks:
        state.bookmarks[name] += 1
        name = f"{name}_{state.bookmarks[name]}"
    else:
        state.bookmarks[name] = 0
    _add_bookmark(state, r._r, name)
    return p


def _next_table_label(state, text):
    """表格编号：正文区 Table 1、2……；附录内 Table A1、A2……（每附录独立计数）。"""
    if state.appendix:
        c = state.app_counters.setdefault(state.appendix, {"table": 0, "figure": 0})
        c["table"] += 1
        return text or f"Table {state.appendix}{c['table']}"
    state.table_no += 1
    return text or f"Table {state.table_no}"


def _next_figure_label(state, text):
    if state.appendix:
        c = state.app_counters.setdefault(state.appendix, {"table": 0, "figure": 0})
        c["figure"] += 1
        return text or f"Figure {state.appendix}{c['figure']}"
    state.figure_no += 1
    return text or f"Figure {state.figure_no}"


def _walk_body(doc, state, lines, render=True):
    """遍历正文标记行。render=True 时生成内容；False 时仅登记编号与书签名，
    使 {REF:Table 1}/{EQREF:1} 等前向引用可用。"""
    quote_mode = False
    after_appendix = False
    for raw in lines:
        line = raw.rstrip("\n").rstrip()
        if not line:
            continue
        if line.startswith("[ENDQUOTE]"):
            quote_mode = False
            continue
        if line.startswith("[QUOTE]"):
            quote_mode = True
            state.appendix = None
            continue
        if quote_mode:
            # 块引用（8.27）：整段左缩进 0.5；[P] 后续段首行再缩进 0.5
            if line.startswith("[P]"):
                if render:
                    para(doc, line.split("]", 1)[1].strip(),
                         indent=0.5, left=0.5, state=state)
            elif render:
                para(doc, line, left=0.5, state=state)
            continue
        if line.startswith("[TABLEDATA]"):
            after_appendix = False
            table_id = line.split("]", 1)[1].strip()
            if render:
                if table_id not in state.table_specs:
                    raise ValueError(f"TABLEDATA 未找到表格 ID：{table_id}")
                add_apa_table_spec(doc, state.table_specs[table_id])
            continue
        if line.startswith("[TABLE]"):
            after_appendix = False
            label = _next_table_label(state, line.split("]", 1)[1].strip())
            state.known_bookmarks.add(label.replace(" ", "_"))
            if render:
                _label_para(doc, state, label, bold=True)
            continue
        if line.startswith("[FIGURE]"):
            after_appendix = False
            label = _next_figure_label(state, line.split("]", 1)[1].strip())
            state.known_bookmarks.add(label.replace(" ", "_"))
            if render:
                _label_para(doc, state, label, bold=True)
            continue
        if line.startswith("[CAPTION]"):
            if render:
                text = line.split("]", 1)[1].strip()
                if after_appendix:
                    para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, tc=True)
                    after_appendix = False
                else:
                    para(doc, text, italic=True, tc=True)
            continue
        if line.startswith("[NOTE]"):
            if render:
                p = doc.add_paragraph()
                r = p.add_run("Note. ")
                r.italic = True
                _set_run_font(r)
                add_runs_with_markup(p, line.split("]", 1)[1].strip(),
                                     doc=doc, state=state)
            continue
        if line.startswith("[APPENDIX]"):
            if render:
                doc.add_page_break()
            text = line.split("]", 1)[1].strip() or "Appendix"
            m = re.search(r"Appendix\s+([A-Z])", text, re.I)
            state.appendix = m.group(1).upper() if m else None
            if render:
                para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, tc=True)
            after_appendix = True
            continue
        if line.startswith("[EQ]"):
            after_appendix = False
            state.appendix = None
            state.eq_no += 1
            state.known_bookmarks.add(f"Eq{state.eq_no}")
            if render:
                render_equation(doc, state, line, state.eq_no)
            continue
        if line.startswith("[H") and len(line) > 4 and line[3] == "]":
            after_appendix = False
            state.appendix = None
            if render:
                render_heading(doc, line)
            continue
        after_appendix = False
        state.appendix = None
        if render:
            para(doc, line, indent=0.5, state=state)


def render_body(doc, title, lines, state):
    doc.add_page_break()
    para(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)  # 正文首页重复标题
    _walk_body(doc, state, lines, render=True)


# ---------------- 参考文献排序（9.44–9.49） ----------------

INITIAL_RE = re.compile(r"^[A-Z]\.(\s*-?[A-Z]\.)*$")
DATE_RE = re.compile(
    r"\((\d{4}(?:[a-z])?(?:\s*[-–]\s*(?:\d{4}|present))?"
    r"|n\.d\.(?:-[a-z])?"
    r"|in\s+press(?:-[a-z])?)[^)]*\)\.?",
    re.I)


def _norm(s):
    """只留小写字母数字：自然实现 nothing precedes something（Loft < Loftus、
    Marshall < Marshall-Petrini）。"""
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _person_key(author):
    surname, sep, initials = author.partition(",")
    if sep:
        return (_norm(surname), _norm(initials))
    return (_norm(author), "")


def _parse_authors(author_part):
    """解析作者部分 → 作者键序列。个人作者 = (姓, 首字母) 元组；
    群体作者/无作者标题 = 单元素元组（去前导冠词，9.49）。"""
    part = re.sub(r"\([^)]*\)", " ", author_part)          # 去 (Ed.) (Director) 等角色描述
    part = re.sub(r",\s*(?:Jr\.|Sr\.|II|III|IV)\s*$", "", part, flags=re.I)
    part = part.strip(" .")
    if not part:
        return (("",),)
    # 逗号会把 "Singh, R." 拆成 "Singh" + "R."：纯首字母段并回前一个人
    segs = [s.strip(" ,.") for s in part.replace("&", ",").split(",") if s.strip(" ,.")]
    persons = []
    for s in segs:
        if INITIAL_RE.match(s) and persons:
            persons[-1] = f"{persons[-1]}, {s}"
        else:
            persons.append(s)
    if len(persons) > 1:
        return tuple(_person_key(p) for p in persons)
    k = re.sub(r"^(a|an|the)\s+", "", persons[0], flags=re.I)  # 群体作者/无作者标题
    return ((_norm(k),),)


def _date_info(entry, m):
    """返回（日期核心, 核心在原串中的结束位置）。"""
    g = m.group(1)
    if re.match(r"n\.d\.", g, re.I):
        return "n.d.", m.start(1) + 4
    if re.match(r"in\s+press", g, re.I):
        return "in press", m.start(1) + 8
    m2 = re.match(r"\d{4}", g)
    if m2:
        return m2.group(0), m.start(1) + 4
    return g, m.end(1)


def _year_key(core):
    c = core.lower()
    if c == "n.d.":
        return (0, 0)  # n.d. 最前
    if c == "in press":
        return (2, 0)  # in press 最后
    m = re.match(r"(\d{4})", c)
    if m:
        return (1, int(m.group(1)))
    return (1, 0)


def _title_key(rest):
    t = re.sub(r"^(a|an|the)\s+", "", rest.strip(), flags=re.I)
    return _norm(t)


def _parse_entry(entry):
    """解析参考文献条目 → (作者键序列, 年份键, 标题键)；无法解析返回 None。"""
    m = DATE_RE.search(entry)
    if not m:
        return None
    authors = _parse_authors(entry[:m.start()])
    core, _ = _date_info(entry, m)
    return (authors, _year_key(core), _title_key(entry[m.end():]))


def _assign_letter(entry, m, letter):
    """在日期核心后插入 a/b/c（n.d./in press 用 -a 形式），替换条目中原有的字母。"""
    core, pos = _date_info(entry, m)
    sep = "-" if core in ("n.d.", "in press") else ""
    suffix = re.sub(r"^-?[a-z](?=$|[\s,-])", "", m.group(1)[len(core):])
    new_g = core + sep + letter + suffix
    return entry[:m.start(1)] + new_g + entry[m.end(1):]


def _sort_references(entries):
    """APA 9.44–9.49 完整排序：作者序列（相同前缀更短者在前 → 单作者优先）→
    年份（n.d. 最前、in press 最后）→ 标题首实义词 → 原始行号（稳定）。
    排序后对「同一第一作者 + 同年」的条目加 a/b/c 后缀（8.18 et al. 歧义场景）。
    返回 (排序后条目列表, 字母赋值提示列表)。"""
    parsed = [(_parse_entry(e), i, e) for i, e in enumerate(entries)]

    def sk(item):
        p, i, e = item
        if p is None:
            return (2, (), (), i)  # 解析失败的排最后，保持原序
        return (0, p[0], p[1], p[2], i)

    ordered = sorted(parsed, key=sk)
    counts = {}
    for p, i, e in ordered:
        if p is None:
            continue
        k = (p[0][0], p[1][0], p[1][1])  # 第一作者键 + 年份
        counts[k] = counts.get(k, 0) + 1
    seen = {}
    out, notes = [], []
    for p, i, e in ordered:
        if p is not None:
            k = (p[0][0], p[1][0], p[1][1])
            if counts[k] >= 2:
                seen[k] = seen.get(k, 0) + 1
                letter = chr(ord("a") + seen[k] - 1)
                m = DATE_RE.search(e)
                e = _assign_letter(e, m, letter)
                notes.append(f"{letter} → {e[:70]}{'…' if len(e) > 70 else ''}")
        out.append(e)
    return out, notes


def render_references(doc, entries, sort_entries=True, state=None):
    entries = [e.rstrip("\n").rstrip() for e in entries if e.strip()]
    if sort_entries:
        ordered, notes = _sort_references(entries)
        if ordered != entries:
            print("提示：参考文献已按 9.44–9.49 完整规则自动排序"
                  "（单作者优先、同作者按日期 n.d. 最前 / in press 最后、"
                  "同年按标题加 a/b/c 后缀）")
        if notes:
            print("提示：已为同年条目加字母后缀，请同步核对文内引用：")
            for n in notes:
                print("  ", n)
        entries = ordered
    doc.add_page_break()
    para(doc, "References", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    for entry in entries:
        para(doc, entry, hanging=True, state=state)  # 悬挂缩进 0.5 英寸


def warn_missing_refs(body_text, ref_text):
    """粗查文内年份与参考文献年份的对应关系，发现不对应时输出警告。"""
    ref_years = set(re.findall(r"\b(?:19|20)\d{2}\b", ref_text))
    body_years = set(re.findall(r"\b(?:19|20)\d{2}\b", body_text))
    missing = sorted(body_years - ref_years)
    if missing:
        print("警告：正文出现以下年份但参考文献中无对应条目：",
              ", ".join(missing), "（请核对文内引用与参考文献是否一一对应）")


def verify(path, expect_running_head=False, expect_footnotes=False):
    """读回校验 APA 7 版式，不满足即抛 AssertionError。返回 Document 供进一步断言。"""
    if expect_footnotes:
        assert "word/footnotes.xml" in zipfile.ZipFile(path).namelist(), "缺少脚注部件"
    d = Document(path)
    st = d.styles["Normal"]
    assert st.font.name == FONT and st.font.size.pt == 12, "Normal 字体应为 TNR 12pt"
    assert st.paragraph_format.line_spacing == 2.0, "Normal 行距应为双倍"
    assert st.element.get_or_add_pPr().find(qn("w:widowControl")) is not None, \
        "Normal 样式缺 widowControl（孤行寡行控制）"
    for si, sec in enumerate(d.sections):
        assert all(m.inches == 1 for m in (sec.top_margin, sec.bottom_margin,
                                           sec.left_margin, sec.right_margin)), \
            f"第 {si} 节页边距应为 1 英寸"
        hdr_xml = sec.header.paragraphs[0]._p.xml
        assert "PAGE" in hdr_xml, f"第 {si} 节页眉缺 PAGE 字段"
        if expect_running_head:
            r0 = sec.header.paragraphs[0].runs[0]
            assert r0.text and len(r0.text) <= 50, "running head 应存在且 ≤50 字符"
    # 逐段：行距双倍（段落未直接设置时从样式解析）、run 字体 TNR 12pt
    for p in d.paragraphs:
        ls = p.paragraph_format.line_spacing
        if ls is None:
            ls = p.style.paragraph_format.line_spacing
        if ls is not None:
            assert ls == 2.0, f"段落行距应为双倍：{p.text[:40]}"
        for r in p.runs:
            if r.font.name is not None:
                assert r.font.name == FONT, f"run 字体应为 TNR：{r.text[:30]}"
            if r.font.size is not None:
                assert r.font.size.pt == 12, f"run 字号应为 12pt：{r.text[:30]}"
    return d


# ---------------- 回归自检 ----------------

def _test_sorting():
    entries = [
        "Singh, R. (2019). Later work. Publisher.",
        "Singh, R., & Shah, S. (2016). Two authors. Publisher.",
        "Singh, R. (n.d.). No date. Publisher.",
        "Singh, R. (n.d.). Another no date. Publisher.",
        "Singh, R. (in press). In press. Publisher.",
        "Singh, R. (2019). Earlier title work. Publisher.",
        "Singh, R. (2015). Single author first. Publisher.",
        "Aarons, B. B., & Zhao, X. (2018). Different author. Publisher.",
        "The early bird catches the worm. (2020). No author. Publisher.",
        "American Psychological Association. (2020). Group author. Publisher.",
        "Marshall, A. (2010). Nothing precedes something. Publisher.",
        "Marshall-Petrini, S. (2011). Hyphen later. Publisher.",
    ]
    ordered, notes = _sort_references(entries)
    expect = [
        "Aarons, B. B., & Zhao, X. (2018). Different author. Publisher.",
        "American Psychological Association. (2020). Group author. Publisher.",
        "The early bird catches the worm. (2020). No author. Publisher.",
        "Marshall, A. (2010). Nothing precedes something. Publisher.",
        "Marshall-Petrini, S. (2011). Hyphen later. Publisher.",
        "Singh, R. (n.d.-a). Another no date. Publisher.",
        "Singh, R. (n.d.-b). No date. Publisher.",
        "Singh, R. (2015). Single author first. Publisher.",
        "Singh, R. (2019a). Earlier title work. Publisher.",
        "Singh, R. (2019b). Later work. Publisher.",
        "Singh, R. (in press). In press. Publisher.",
        "Singh, R., & Shah, S. (2016). Two authors. Publisher.",
    ]
    assert ordered == expect, "参考文献排序错误：\n" + "\n".join(ordered)
    assert len(notes) == 4, "应有两组同年条目各分配 a/b"
    print("参考文献排序测试通过")


def _test_keywords():
    kws, kept = lowercase_keywords(
        "COVID-19 pandemic, United States, Big Five, depression symptoms, "
        "cognitive behavioral therapy, ADHD")
    assert kws == ("COVID-19 pandemic, United States, Big Five, depression symptoms, "
                   "cognitive behavioral therapy, ADHD"), kws
    assert "COVID-19" in kept and "United" in kept and "States" in kept
    kws2, _ = lowercase_keywords("Depression in college students")
    assert kws2 == "depression in college students", kws2
    kws3, _ = lowercase_keywords("United States", protect="United States")
    assert kws3 == "United States", kws3
    kws4, kept4 = lowercase_keywords("kw1, study 1")
    assert kws4 == "kw1, study 1" and not kept4, "小写带数字词不应视为专名"
    print("关键词专名保护测试通过")


def _test_running_head():
    long_title = title_case("a study of the effects of social media on the "
                            "mental health of college students during the pandemic")
    rh = _make_running_head(long_title, "")
    assert len(rh) <= 50 and rh == rh.upper(), "running head 应截断为 ≤50 大写"
    assert _make_running_head("", "CUSTOM HEAD") == "CUSTOM HEAD"
    assert title_case("appendix a") == "Appendix A", "单字母词应保留大写"
    print("running head / title case 测试通过")


def run_selftest():
    """固定样例回归测试：覆盖标题页/摘要/五级标题/块引用/表格图/附录编号/
    脚注/公式/交叉引用/超链接/参考文献排序，并对关键版式做断言。"""
    _test_sorting()
    _test_keywords()
    _test_running_head()

    tmp = tempfile.mkdtemp(prefix="apa7_selftest_")
    path = os.path.join(tmp, "selftest.docx")
    body_txt = "\n".join([
        "[H1]method",
        "[H2]participants",
        "Two hundred participants were recruited.{FN:Sample size was determined by a power analysis.}",
        "[H4]scoring procedure. Scores ranged from 1 to 7.",
        "The model appears in {EQREF:1}, and the results appear in {REF:Table 1}.",
        "[EQ]d = \\frac{\\sum_{i=1}^{n} (x_i - M)^2}{n - 1}",
        "[QUOTE]",
        "Inner speech is a paradoxical phenomenon that presents considerable challenges.",
        "[P]The present research sheds light on this issue.",
        "[ENDQUOTE]",
        "[TABLE]",
        "[CAPTION]descriptive statistics for study measures",
        "[NOTE]General note here.",
        "[NOTE]^a^ Refers to the first row.",
        "[FIGURE]",
        "[CAPTION]bar graph of mean scores",
        "[APPENDIX]Appendix A",
        "[CAPTION]stimulus materials",
        "[TABLE]",
        "[CAPTION]stimulus item statistics",
        "[TABLE]",
        "[CAPTION]more stimulus item statistics",
    ])
    refs_txt = "\n".join([
        "Smith, J. A. (2020). Title of article. *Journal of Psychology*, 10(2), 34-56. https://doi.org/10.1000/abc123",
        "Wang, L. (2019). Title of book. Publisher.",
    ])
    note_txt = "\n".join([
        "Ming Li https://orcid.org/0000-0000-0000-0000",
        "We have no known conflict of interest to disclose.",
    ])
    body_file = os.path.join(tmp, "body.txt")
    refs_file = os.path.join(tmp, "refs.txt")
    note_file = os.path.join(tmp, "note.txt")
    abs_file = os.path.join(tmp, "abstract.txt")
    with open(body_file, "w", encoding="utf-8") as f:
        f.write(body_txt)
    with open(refs_file, "w", encoding="utf-8") as f:
        f.write(refs_txt)
    with open(note_file, "w", encoding="utf-8") as f:
        f.write(note_txt)
    with open(abs_file, "w", encoding="utf-8") as f:
        f.write("This study examined the effect of X on Y among 200 participants.")

    doc = Document()
    set_normal_style(doc)
    set_margins(doc)
    title = title_case("the effect of X on Y")
    assert title == "The Effect of X on Y", "title case 转换错误"
    assert title_case("中文标题测试") == "中文标题测试", "中文标题不应被转换"
    assert title_case("the effects of") == "The Effects of", "末词短介词不应大写"
    assert title_case("What if we sleep?") == "What if We Sleep?", "if 应保持小写"
    running = "THE EFFECT OF X ON Y"
    setup_header(doc, running)
    title_page(doc, argparse.Namespace(
        mode="professional", title=title,
        author="Ming Li^1^; Zhang San^2^",
        affiliation="Department of Psychology, XX University^1^; School of Education, YY University^2^",
        course="", instructor="", date="", author_note=note_file), title)
    abstract_page(doc, argparse.Namespace(abstract=abs_file, keywords="kw1, kw2"))
    # 超长摘要（>250 词）应输出警告
    long_file = os.path.join(tmp, "abstract_long.txt")
    with open(long_file, "w", encoding="utf-8") as f:
        f.write(" ".join(["word"] * 251))
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        abstract_page(Document(), argparse.Namespace(
            abstract=long_file, keywords=""))
    assert "250" in buf.getvalue(), "超长摘要应输出警告"
    scan = DocState()
    _walk_body(None, scan, body_txt.splitlines(), render=False)  # 预扫描：前向引用
    state = DocState()
    state.known_bookmarks = scan.known_bookmarks
    render_body(doc, title, body_txt.splitlines(), state)
    render_references(doc, refs_txt.splitlines(), state=state)
    doc.save(path)
    d = verify(path, expect_running_head=True, expect_footnotes=True)

    ps = d.paragraphs
    # 摘要：标签加粗居中、正文单段首行不缩进；关键词行斜体标签 + 小写关键词
    abs_p = next(p for p in ps if p.text.startswith("This study examined"))
    fli = abs_p.paragraph_format.first_line_indent
    assert fli is None or fli == 0, "摘要首行不应缩进"
    abs_label = next(p for p in ps if p.text == "Abstract")
    assert abs_label.runs[0].bold and abs_label.alignment == WD_ALIGN_PARAGRAPH.CENTER
    kw_p = next(p for p in ps if p.text.startswith("Keywords:"))
    assert kw_p.runs[0].italic and kw_p.text == "Keywords: kw1, kw2", "关键词行格式"
    # 分页结构：摘要、正文、附录、参考文献各一个显式分页符
    assert sum(1 for p in ps if 'w:type="page"' in p._p.xml) == 4, "分页符数量"
    # H4：仅标题 run 加粗，其后正文不加粗
    h4 = next(p for p in ps if p.runs and p.runs[0].text.startswith("Scoring Procedure"))
    assert h4.runs[0].bold is True, "H4 标题应加粗"
    assert all(r.bold in (False, None) for r in h4.runs[1:]), "H4 同行正文不应加粗"
    # 块引用：左缩进 0.5、首行不缩进；[P] 段首行再缩进 0.5
    quote = next(p for p in ps if "Inner speech" in p.text)
    assert quote.paragraph_format.left_indent.inches == 0.5, "块引用左缩进"
    assert not quote.paragraph_format.first_line_indent, "块引用首行不应缩进"
    quote2 = next(p for p in ps if "The present research" in p.text)
    assert quote2.paragraph_format.first_line_indent.inches == 0.5, "块引用第二段首行缩进"
    # 多作者 and 连接、多单位各占一行、作者右上标
    assert any("Ming Li" in p.text and "and Zhang San" in p.text
               for p in ps), "多作者应 and 连接"
    assert any("School of Education, YY University" in p.text
               for p in ps), "多单位各占一行"
    assert any(r.font.superscript for p in ps for r in p.runs), "作者编号应上标"
    # 参考文献：悬挂缩进、斜体刊名、DOI 超链接
    ref = next(p for p in ps if p.text.startswith("Smith"))
    assert ref.paragraph_format.left_indent.inches == 0.5
    assert ref.paragraph_format.first_line_indent.inches == -0.5
    assert any(r.italic for r in ref.runs), "参考文献斜体标记应生效"
    assert 'hyperlink' in ref._p.xml and 'r:id' in ref._p.xml, "DOI 应转超链接"
    ext = [r for r in d.part.rels.values()
           if r.is_external and r.reltype == RT.HYPERLINK]
    assert any(r.target_ref.startswith("https://doi.org/") for r in ext), "外部超链接关系"
    # 作者注正文：左对齐、首行缩进 0.5（2.7）
    note_p = next(p for p in ps if "no known conflict" in p.text)
    assert note_p.alignment == WD_ALIGN_PARAGRAPH.LEFT, "作者注正文应左对齐"
    assert note_p.paragraph_format.first_line_indent.inches == 0.5, "作者注正文应首行缩进"
    # 特定注：上标小写字母（7.14）
    sp = next(p for p in ps if "Refers to the first row" in p.text)
    assert any(r.font.superscript for r in sp.runs), "特定注字母应上标"
    # 表图编号与附录编号（7.2 / 2.14：附录内 Table A1、A2）
    t1 = next(p for p in ps if p.text == "Table 1")
    assert t1.runs[0].bold, "Table 1 编号应加粗"
    assert any(p.text == "Table A1" for p in ps), "附录表应编号 Table A1"
    assert any(p.text == "Table A2" for p in ps), "附录第二张表应编号 Table A2"
    cap = next(p for p in ps if p.text.startswith("Descriptive Statistics"))
    assert cap.runs[0].italic, "题注应斜体"
    note = next(p for p in ps if p.text.startswith("Note."))
    assert note.runs[0].italic, "表注 Note. 应斜体"
    app = next(p for p in ps if p.text == "Appendix A")
    assert app.runs[0].bold, "附录标签应加粗"
    # 书签与交叉引用
    assert any('bookmarkStart' in p._p.xml and 'Table_A1' in p._p.xml
               for p in ps), "附录表编号应有书签"
    assert any('REF Table_1' in p._p.xml for p in ps), "题注交叉引用 REF 域"
    assert any('REF Eq1' in p._p.xml for p in ps), "公式交叉引用 REF 域"
    # 公式：m:oMath 对象 + 编号书签 + 编号文本
    assert any('oMath' in p._p.xml for p in ps), "公式应为 m:oMath 对象"
    assert any('bookmarkStart' in p._p.xml and 'Eq1' in p._p.xml
               for p in ps), "公式编号应有书签"
    assert any(r.text == "(1)" for p in ps for r in p.runs), "公式编号 (1)"
    # 脚注：部件、内容、正文引用
    fn_part = next(r.target_part for r in d.part.rels.values()
                   if r.reltype == RT.FOOTNOTES)
    fn_elm = parse_xml(fn_part.blob)
    fns = fn_elm.findall(qn("w:footnote"))
    assert len(fns) == 3, "2 个特殊脚注 + 1 条内容脚注"
    assert fns[0].get(qn("w:id")) == "-1", "separator 脚注 id 应为 -1"
    from lxml import etree
    assert "power analysis" in etree.tostring(fn_elm, encoding="unicode"), \
        "脚注文本应写入 footnotes.xml"
    assert any("footnoteReference" in p._p.xml for p in ps), "正文应有脚注引用"
    assert d.settings.element.find(qn("w:footnotePr")) is not None, "settings 缺 footnotePr"
    assert d.settings.element.find(qn("w:updateFields")) is not None, "settings 缺 updateFields"
    # running head 与页码同一段
    hdr = d.sections[0].header.paragraphs[0]
    assert hdr.runs[0].text == running and "PAGE" in hdr._p.xml, "running head 与页码应同行"
    print("SELFTEST PASS (样例输出:", path, ")")
    return True


def main():
    ap = argparse.ArgumentParser(description="生成 APA 7 论文 .docx")
    ap.add_argument("--mode", choices=["student", "professional"],
                    default="student", help="学生论文（默认）/专业论文")
    ap.add_argument("--title", default="", help="论文标题（英文自动转 title case）")
    ap.add_argument("--author", default="",
                    help="作者姓名，; 分隔多作者（支持 ^1^ 上标编号）")
    ap.add_argument("--affiliation", default="",
                    help="院系/单位，; 分隔多条（各条带 ^1^ 编号与作者对应）")
    ap.add_argument("--course", default="", help="课程编号与名称（学生版）")
    ap.add_argument("--instructor", default="", help="教师姓名（学生版）")
    ap.add_argument("--date", default="", help="截止日期（学生版）")
    ap.add_argument("--author-note", default="", help="作者注文本文件，每行一段（专业版）")
    ap.add_argument("--running-head", default="",
                    help="running head（专业版；缺省由标题自动生成：大写、& 代替 and、截 50 字符）")
    ap.add_argument("--abstract", default="", help="摘要文本文件（生成摘要页）")
    ap.add_argument("--keywords", default="", help="关键词，逗号分隔（摘要页 Keywords 行）")
    ap.add_argument("--keywords-protect", default="",
                    help="关键词专有名词白名单，逗号分隔（整条保留大小写）")
    ap.add_argument("--body", default="", help="正文文本文件（标记语法见文件头 docstring）")
    ap.add_argument("--tables", default="",
                    help="JSON 表格规格文件；正文用 [TABLEDATA]表格ID 插入")
    ap.add_argument("--references", default="", help="参考文献文本文件（每行一条，支持行内标记）")
    ap.add_argument("--no-sort", action="store_true",
                    help="参考文献保持输入顺序，不自动按 APA 规则重排")
    ap.add_argument("--output", default="paper.docx", help="输出路径")
    ap.add_argument("--selftest", action="store_true", help="运行回归自检后退出")
    args = ap.parse_args()

    if args.selftest:
        sys.exit(0 if run_selftest() else 1)

    if not args.title:
        ap.error("--title 必填")
    title = title_case(args.title)

    doc = Document()
    set_normal_style(doc)
    set_margins(doc)
    if args.mode == "professional":
        if args.running_head and len(args.running_head) > 50:
            print(f"警告：running head 超过 50 字符（{len(args.running_head)}），"
                  "APA 2.8 要求 ≤50 字符")
        running = _make_running_head(title, args.running_head)
    else:
        running = ""
    setup_header(doc, running)

    title_page(doc, args, title)

    if args.abstract or args.keywords:
        abstract_page(doc, args)

    body_lines = open(args.body, encoding="utf-8").readlines() if args.body else []
    scan = DocState()
    _walk_body(None, scan, body_lines, render=False)  # 预扫描编号：支持前向引用
    state = DocState()
    state.known_bookmarks = scan.known_bookmarks
    if args.tables:
        with open(args.tables, encoding="utf-8") as table_file:
            state.table_specs = json.load(table_file)
    render_body(doc, title, body_lines, state)

    ref_lines = open(args.references, encoding="utf-8").readlines() if args.references else []
    render_references(doc, ref_lines, sort_entries=not args.no_sort, state=state)
    if body_lines and ref_lines:
        warn_missing_refs("\n".join(body_lines), "\n".join(ref_lines))

    doc.save(args.output)
    verify(args.output, expect_running_head=bool(running),
           expect_footnotes=state.footnotes_part is not None)
    print("已生成:", args.output, "| 版式校验通过")


if __name__ == "__main__":
    main()
