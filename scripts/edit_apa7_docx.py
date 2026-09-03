# -*- coding: utf-8 -*-
"""基于 python-docx 的 APA 7 既有文档编辑器。"""

import argparse
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches

from apa7_word_components import (
    add_apa_table_spec,
    add_table_caption,
    add_table_note,
    append_stat_runs,
    apply_apa7_layout,
)


ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _style_level(paragraph):
    style_id = paragraph.style.style_id if paragraph.style else ""
    match = re.search(r"Heading(\d+)$", style_id, re.I)
    if not match and paragraph.style:
        match = re.search(r"Heading\s+(\d+)$", paragraph.style.name, re.I)
    return int(match.group(1)) if match else None


def _element_heading_level(element):
    if element.tag != qn("w:p"):
        return None
    p_style = element.find("./w:pPr/w:pStyle", namespaces=element.nsmap)
    if p_style is None:
        return None
    value = p_style.get(qn("w:val"), "")
    match = re.search(r"Heading(\d+)$", value, re.I)
    return int(match.group(1)) if match else None


def _find_paragraph(doc, text, *, contains=False):
    matches = [p for p in doc.paragraphs
               if (text in p.text if contains else p.text.strip() == text.strip())]
    if len(matches) != 1:
        raise ValueError(f"段落定位应唯一，实际找到 {len(matches)} 处：{text}")
    return matches[0]


def _clear_paragraph_text(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _render_content(doc, content):
    for block in content:
        block_type = block.get("type")
        if block_type == "paragraph":
            paragraph = doc.add_paragraph(style=block.get("style"))
            paragraph.alignment = ALIGNMENTS.get(
                block.get("alignment", "left"), WD_ALIGN_PARAGRAPH.LEFT)
            indent = block.get("first_line_indent", 0.5)
            paragraph.paragraph_format.first_line_indent = Inches(indent)
            append_stat_runs(
                paragraph, block.get("text", ""),
                bold=block.get("bold", False),
                italic=block.get("italic", False),
                size=block.get("font_size", 12),
            )
        elif block_type == "heading":
            level = int(block.get("level", 1))
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            append_stat_runs(
                paragraph, block.get("text", ""), bold=True,
                italic=level in {3, 5})
        elif block_type == "table":
            if block.get("number") or block.get("title"):
                add_table_caption(
                    doc, block.get("number", "Table"), block.get("title", ""))
            add_apa_table_spec(doc, block)
            if block.get("note"):
                add_table_note(doc, block["note"])
        elif block_type == "table_note":
            add_table_note(doc, block.get("text", ""), block.get("font_size", 10))
        else:
            raise ValueError(f"不支持的内容块类型：{block_type}")


def _collect_new_body_elements(doc, before):
    sect_pr = doc._element.body.sectPr
    return [element for element in doc._element.body
            if element is not sect_pr and element not in before]


def _render_and_move_after(doc, anchor, content):
    before = set(doc._element.body)
    _render_content(doc, content)
    new_elements = _collect_new_body_elements(doc, before)
    current = anchor
    for element in new_elements:
        current.addnext(element)
        current = element


def _replace_section(doc, operation):
    heading = _find_paragraph(
        doc, operation["heading"], contains=operation.get("contains", False))
    level = _style_level(heading)
    if level is None:
        raise ValueError(f"目标段落不是 Word 标题样式：{heading.text}")
    anchor = heading._p
    current = anchor.getnext()
    while current is not None:
        next_element = current.getnext()
        next_level = _element_heading_level(current)
        if next_level is not None and next_level <= level:
            break
        doc._element.body.remove(current)
        current = next_element
    _render_and_move_after(doc, anchor, operation.get("content", []))


def _replace_paragraph(doc, operation):
    paragraph = _find_paragraph(
        doc, operation["match"], contains=operation.get("contains", False))
    _clear_paragraph_text(paragraph)
    append_stat_runs(paragraph, operation.get("text", ""))


def _insert_after(doc, operation):
    paragraph = _find_paragraph(
        doc, operation["match"], contains=operation.get("contains", False))
    _render_and_move_after(doc, paragraph._p, operation.get("content", []))


def edit_document(input_path, output_path, spec):
    """按结构化操作修改既有 DOCX，并始终另存为新文件。"""
    input_path = Path(input_path)
    output_path = Path(output_path)
    if input_path.resolve() == output_path.resolve():
        raise ValueError("output_path 必须与 input_path 不同，避免覆盖原稿")
    if isinstance(spec, (str, Path)):
        spec = json.loads(Path(spec).read_text(encoding="utf-8"))

    doc = Document(input_path)
    if spec.get("normalize_apa", False):
        apply_apa7_layout(doc, normalize_runs=True)
    for operation in spec.get("operations", []):
        op_type = operation.get("type")
        if op_type == "replace_section":
            _replace_section(doc, operation)
        elif op_type == "replace_paragraph":
            _replace_paragraph(doc, operation)
        elif op_type == "insert_after":
            _insert_after(doc, operation)
        elif op_type == "append_content":
            _render_content(doc, operation.get("content", []))
        else:
            raise ValueError(f"不支持的操作类型：{op_type}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    Document(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="保格式修改既有 APA 7 DOCX")
    parser.add_argument("--input", required=True, help="输入 DOCX")
    parser.add_argument("--output", required=True, help="输出 DOCX；不得覆盖输入文件")
    parser.add_argument("--spec", required=True, help="JSON 编辑操作文件")
    args = parser.parse_args()
    path = edit_document(args.input, args.output, args.spec)
    print("已生成：", path)


if __name__ == "__main__":
    main()
