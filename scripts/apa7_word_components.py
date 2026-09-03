# -*- coding: utf-8 -*-
"""可复用的 APA 7 Word 组件：版式、统计格式与表格。"""

from decimal import Decimal, ROUND_HALF_UP
import re

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


FONT = "Times New Roman"
FONT_EAST = "宋体"
STAT_RE = re.compile(
    r"(?<!top-)(?<![A-Za-z])"
    r"(?:n(?=\d)|(Mdn|SD|SE|OR|Pr|df|R²|N|n|M|F|t|d|p|r)"
    r"(?=(?:s\b|\b)))"
)


def _decimal(value):
    return value if isinstance(value, Decimal) else Decimal(str(value))


def _round(value, digits):
    quantum = Decimal("1").scaleb(-digits)
    return _decimal(value).quantize(quantum, rounding=ROUND_HALF_UP)


def format_probability(value, digits=3):
    """格式化概率或比例；按 APA 规则省略前导零。"""
    return f"{_round(value, digits):.{digits}f}".lstrip("0")


def format_p(value, digits=3, threshold=0.001):
    """格式化数值 p；小于阈值时返回 ``< .001``。"""
    value = _decimal(value)
    if value < _decimal(threshold):
        return f"< {format_probability(threshold, digits)}"
    return format_probability(value, digits)


def format_or(value, digits=2):
    """格式化优势比；OR 可能大于 1，因此保留前导零。"""
    return f"{_round(value, digits):.{digits}f}"


def format_ci(low, high, digits=2, leading_zero=True):
    if leading_zero:
        left, right = format_or(low, digits), format_or(high, digits)
    else:
        left, right = format_probability(low, digits), format_probability(high, digits)
    return f"[{left}, {right}]"


def format_or_ci(estimate, low, high, digits=2):
    return f"{format_or(estimate, digits)} {format_ci(low, high, digits, True)}"


def _set_run_font(run, size=12):
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), FONT)
    rpr.rFonts.set(qn("w:hAnsi"), FONT)
    rpr.rFonts.set(qn("w:eastAsia"), FONT_EAST)


def _add_run(paragraph, text, *, bold=False, italic=False, subscript=False,
             size=12):
    if not text:
        return None
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.subscript = subscript
    _set_run_font(run, size)
    return run


def iter_stat_segments(text):
    """产出 ``(文本, 是否斜体, 是否下标)``，供不同渲染器共享。"""
    cursor = 0
    for match in STAT_RE.finditer(text):
        if match.start() > cursor:
            yield text[cursor:match.start()], False, False
        token = match.group(0)
        yield token, True, False
        cursor = match.end()
        if token == "n":
            digits = re.match(r"\d+", text[cursor:])
            if digits:
                value = digits.group(0)
                yield value, False, True
                cursor += len(value)
    if cursor < len(text):
        yield text[cursor:], False, False


def append_stat_runs(paragraph, text, *, bold=False, italic=False, size=12):
    """追加文本，并自动排版常见统计符号及 n0/n1 的下标。"""
    for value, stat_italic, subscript in iter_stat_segments(text):
        _add_run(paragraph, value, bold=bold,
                 italic=italic or stat_italic,
                 subscript=subscript, size=size)
    return paragraph


def _set_style_font(style, size=12, bold=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EAST)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def apply_apa7_layout(doc, *, normalize_runs=False):
    """应用 US Letter、1 英寸页边距和 APA 7 常用段落样式。"""
    normal = doc.styles["Normal"]
    _set_style_font(normal, 12)
    pf = normal.paragraph_format
    pf.line_spacing = 2.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0.5)
    pf.widow_control = True

    for level in range(1, 6):
        name = f"Heading {level}"
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        _set_style_font(style, 12, bold=True, italic=level in {3, 5})
        hpf = style.paragraph_format
        hpf.line_spacing = 2.0
        hpf.space_before = Pt(0)
        hpf.space_after = Pt(0)
        hpf.widow_control = True
        hpf.keep_with_next = True
        hpf.first_line_indent = Inches(0.5 if level in {4, 5} else 0)
        hpf.alignment = (WD_ALIGN_PARAGRAPH.CENTER if level == 1
                         else WD_ALIGN_PARAGRAPH.LEFT)

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    if normalize_runs:
        for paragraph in doc.paragraphs:
            paragraph.paragraph_format.widow_control = True
            for run in paragraph.runs:
                _set_run_font(run, 12)
    return doc


def _set_on_off(parent, tag, value="true"):
    element = parent.find(qn(f"w:{tag}"))
    if element is None:
        element = OxmlElement(f"w:{tag}")
        parent.append(element)
    element.set(qn("w:val"), value)
    return element


def _set_cell_width(cell, dxa):
    cell.width = Inches(dxa / 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def _set_cell_margins(cell, top=70, bottom=70, start=90, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("bottom", bottom),
                        ("start", start), ("end", end)):
        elm = margins.find(qn(f"w:{edge}"))
        if elm is None:
            elm = OxmlElement(f"w:{edge}")
            margins.append(elm)
        elm.set(qn("w:w"), str(value))
        elm.set(qn("w:type"), "dxa")


def _set_table_properties(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge, value, size in (
        ("top", "single", "8"), ("bottom", "single", "8"),
        ("start", "nil", "0"), ("end", "nil", "0"),
        ("insideH", "nil", "0"), ("insideV", "nil", "0"),
    ):
        elm = borders.find(qn(f"w:{edge}"))
        if elm is None:
            elm = OxmlElement(f"w:{edge}")
            borders.append(elm)
        elm.set(qn("w:val"), value)
        elm.set(qn("w:sz"), size)
        elm.set(qn("w:color"), "000000")


def _set_header_bottom_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "000000")


def _prepare_row(row, repeat=False):
    tr_pr = row._tr.get_or_add_trPr()
    _set_on_off(tr_pr, "cantSplit")
    if repeat:
        _set_on_off(tr_pr, "tblHeader")


def _fill_cell(cell, value, *, width_dxa, header=False, first=False,
               font_size=9):
    _set_cell_width(cell, width_dxa)
    _set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if header or not first
                           else WD_ALIGN_PARAGRAPH.LEFT)
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.widow_control = True
    append_stat_runs(paragraph, str(value), bold=header, size=font_size)
    if header:
        _set_header_bottom_border(cell)


def add_apa_table(doc, headers, rows, *, widths=None, font_size=9,
                  first_column_left=True):
    """添加具有固定列宽、重复表头和 APA 横线的 Word 表格。"""
    if not headers:
        raise ValueError("headers 不能为空")
    if any(len(row) != len(headers) for row in rows):
        raise ValueError("每行单元格数必须与 headers 一致")
    if not 8 <= float(font_size) <= 12:
        raise ValueError("表格字号应在 8–12 pt 之间")

    section = doc.sections[-1]
    content_width = (section.page_width - section.left_margin -
                     section.right_margin) / 914400
    if widths is None:
        widths = [content_width / len(headers)] * len(headers)
    if len(widths) != len(headers) or any(float(w) <= 0 for w in widths):
        raise ValueError("widths 必须与 headers 等长且均为正数")
    scale = content_width / sum(float(w) for w in widths)
    widths_dxa = [round(float(w) * scale * 1440) for w in widths]
    widths_dxa[-1] += round(content_width * 1440) - sum(widths_dxa)

    table = doc.add_table(rows=1, cols=len(headers))
    _set_table_properties(table, widths_dxa)
    _prepare_row(table.rows[0], repeat=True)
    for index, value in enumerate(headers):
        _fill_cell(table.rows[0].cells[index], value,
                   width_dxa=widths_dxa[index], header=True,
                   font_size=font_size)
    for values in rows:
        row = table.add_row()
        _prepare_row(row)
        for index, value in enumerate(values):
            _fill_cell(row.cells[index], value,
                       width_dxa=widths_dxa[index],
                       first=first_column_left and index == 0,
                       font_size=font_size)
    return table


def add_table_caption(doc, number, title):
    number_p = doc.add_paragraph()
    number_p.paragraph_format.first_line_indent = Inches(0)
    number_p.paragraph_format.keep_with_next = True
    append_stat_runs(number_p, number, bold=True)
    title_p = doc.add_paragraph()
    title_p.paragraph_format.first_line_indent = Inches(0)
    title_p.paragraph_format.keep_with_next = True
    append_stat_runs(title_p, title, italic=True)
    return number_p, title_p


def add_table_note(doc, text, size=10):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Inches(0)
    _add_run(paragraph, "Note. ", italic=True, size=size)
    append_stat_runs(paragraph, text, size=size)
    return paragraph


def add_apa_table_spec(doc, spec):
    """从 JSON 兼容字典添加单表或多面板表格。"""
    headers = spec.get("headers")
    widths = spec.get("widths")
    font_size = spec.get("font_size", 9)
    panels = spec.get("panels")
    result = []
    if panels:
        for panel in panels:
            label = panel.get("label")
            if label:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.keep_with_next = True
                append_stat_runs(p, label, italic=True)
            result.append(add_apa_table(
                doc, headers, panel.get("rows", []), widths=widths,
                font_size=font_size,
                first_column_left=spec.get("first_column_left", True),
            ))
    else:
        result.append(add_apa_table(
            doc, headers, spec.get("rows", []), widths=widths,
            font_size=font_size,
            first_column_left=spec.get("first_column_left", True),
        ))
    return result


def iter_all_paragraphs(doc, include_headers=True):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    if include_headers:
        for section in doc.sections:
            yield from section.header.paragraphs
            yield from section.footer.paragraphs
